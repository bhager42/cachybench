"""Build the Excel workbook from a results session.

Layout:
  Summary       one row per target, median across runs, plus a comparison chart
  Environment   the machine state the session ran under
  <Target>      per-run metrics table, a bar chart of the FPS figures, and a
                per-second FPS timeline with a line chart
  Synthetics    scores from command-line benchmarks

Charts are native Excel charts pointed at real cells, so editing or extending
the data updates them — nothing is a pasted image.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .metrics import METRIC_LABELS

FONT = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
SECTION_FONT = Font(name=FONT, bold=True, size=11, color="1F3864")
BODY_FONT = Font(name=FONT, size=10)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="666666")
BAND_FILL = PatternFill("solid", fgColor="F2F5FA")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Metrics shown on every per-target sheet, in this order.
HEADLINE_METRICS = [
    "avg_fps",
    "fps_1_low",
    "fps_0p1_low",
    "fps_p99",
    "fps_min",
    "fps_max",
    "frametime_avg_ms",
    "frametime_p99_ms",
    "stutter_pct",
]

CHARTED_METRICS = ["avg_fps", "fps_1_low", "fps_0p1_low"]

# Everything else worth reporting, if MangoHud captured it.
TELEMETRY_METRICS = [
    "gpu_load_avg",
    "gpu_temp_avg",
    "gpu_temp_max",
    "gpu_junction_temp_max",
    "gpu_power_avg",
    "gpu_core_clock_avg",
    "gpu_mem_clock_avg",
    "gpu_vram_used_max",
    "cpu_load_avg",
    "cpu_temp_avg",
    "cpu_temp_max",
    "cpu_power_avg",
    "cpu_mhz_avg",
    "ram_used_max",
]


def _sanitize_sheet_name(name: str, used: set[str]) -> str:
    """Excel sheet names: 31 chars, no []:*?/\\ , and unique."""
    clean = "".join("-" if c in "[]:*?/\\" else c for c in name).strip() or "Sheet"
    clean = clean[:31]
    if clean not in used:
        used.add(clean)
        return clean
    for i in range(2, 100):
        suffix = f" ({i})"
        candidate = clean[: 31 - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
    raise ValueError(f"Could not make a unique sheet name from {name!r}")


def _label(key: str) -> str:
    return METRIC_LABELS.get(key, key.replace("_", " ").title())


def _write_header_row(ws: Worksheet, row: int, values: list[str], start_col: int = 1) -> None:
    for offset, value in enumerate(values):
        cell = ws.cell(row=row, column=start_col + offset, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _style_body(ws: Worksheet, first_row: int, last_row: int, first_col: int, last_col: int) -> None:
    for r in range(first_row, last_row + 1):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
            if (r - first_row) % 2 == 1:
                cell.fill = BAND_FILL


def _autosize(ws: Worksheet, max_col: int, max_width: int = 34) -> None:
    for col in range(1, max_col + 1):
        longest = 0
        for cell in ws[get_column_letter(col)]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max(11, longest + 2), max_width)


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------


def _build_summary(ws: Worksheet, session: dict[str, Any]) -> None:
    ws["A1"] = "Benchmark Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Session {session.get('session_id', '')} — "
        f"{session.get('started', '')} — median of all runs per target"
    )
    ws["A2"].font = NOTE_FONT

    game_results = [r for r in session["results"] if r.get("kind") == "game" and r.get("runs")]

    header_row = 4
    columns = ["Target", "Runs", "Run Spread %", *[_label(m) for m in HEADLINE_METRICS]]
    telemetry_present = [
        m
        for m in TELEMETRY_METRICS
        if any(r["aggregate"].get(m) is not None for r in game_results)
    ]
    columns += [_label(m) for m in telemetry_present]
    _write_header_row(ws, header_row, columns)

    row = header_row + 1
    for result in game_results:
        agg = result["aggregate"]
        ws.cell(row=row, column=1, value=result["name"])
        ws.cell(row=row, column=2, value=agg.get("run_count"))
        ws.cell(row=row, column=3, value=agg.get("run_spread_pct"))
        for i, metric in enumerate(HEADLINE_METRICS):
            ws.cell(row=row, column=4 + i, value=agg.get(metric))
        for i, metric in enumerate(telemetry_present):
            ws.cell(row=row, column=4 + len(HEADLINE_METRICS) + i, value=agg.get(metric))
        row += 1

    last_row = row - 1
    if last_row >= header_row + 1:
        _style_body(ws, header_row + 1, last_row, 1, len(columns))
        ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "FPS comparison across targets (median of runs)"
        chart.y_axis.title = "FPS"
        chart.x_axis.title = "Target"
        chart.height = 9
        chart.width = 22

        # Columns 4, 5, 6 are avg / 1% low / 0.1% low.
        data = Reference(ws, min_col=4, max_col=6, min_row=header_row, max_row=last_row)
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"A{last_row + 3}")

    _autosize(ws, len(columns))

    note_row = last_row + 22
    ws.cell(row=note_row, column=1, value="How to read these numbers").font = SECTION_FONT
    notes = [
        "Avg FPS is derived from mean frametime (1000 / mean ms), not from averaging the "
        "per-frame FPS column — averaging FPS over-weights fast frames.",
        "1% Low / 0.1% Low are the average FPS across the slowest 1% and 0.1% of frames. "
        "They describe how bad the rough moments get, and matter more than the average for "
        "how a game actually feels.",
        "Stutter % is the share of frames taking more than twice the median frametime.",
        "Run Spread % is the gap between the best and worst run as a share of the mean. "
        "Above roughly 3% the run-to-run noise is large enough that small differences "
        "between targets are not meaningful.",
    ]
    for i, text in enumerate(notes):
        cell = ws.cell(row=note_row + 1 + i, column=1, value="• " + text)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=note_row + 1 + i, start_column=1, end_row=note_row + 1 + i, end_column=8)
        ws.row_dimensions[note_row + 1 + i].height = 28


def _build_environment(ws: Worksheet, session: dict[str, Any]) -> None:
    ws["A1"] = "Test Environment"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Captured at the start of the session. Compare this between workbooks before comparing scores."
    ws["A2"].font = NOTE_FONT

    _write_header_row(ws, 4, ["Property", "Value"])
    row = 5
    for key, value in session.get("environment", {}).items():
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=str(value))
        row += 1

    for key in ("session_id", "started", "finished", "config_file", "cachybench_version"):
        if key in session:
            ws.cell(row=row, column=1, value=key.replace("_", " ").title())
            ws.cell(row=row, column=2, value=str(session[key]))
            row += 1

    _style_body(ws, 5, row - 1, 1, 2)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90

    issues = session.get("warnings", [])
    if issues:
        row += 2
        ws.cell(row=row, column=1, value="Warnings raised at session start").font = SECTION_FONT
        for i, text in enumerate(issues, start=1):
            cell = ws.cell(row=row + i, column=1, value="• " + text)
            cell.font = NOTE_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _build_game_sheet(ws: Worksheet, result: dict[str, Any]) -> None:
    ws["A1"] = result["name"]
    ws["A1"].font = TITLE_FONT
    subtitle = f"App ID {result.get('appid', '?')} · mode: {result.get('mode', '?')}"
    if result.get("notes"):
        subtitle += f" · {result['notes']}"
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT

    # Declared game settings, so two workbooks can be compared on more than
    # the game's name. Anything the game reset is called out inline.
    settings = result.get("settings") or {}
    if settings.get("values"):
        drift = settings.get("drift") or {}
        parts = []
        for name, value in settings["values"].items():
            short = name.rsplit("/", 1)[-1]
            if name in drift:
                parts.append(f"{short}={value!r} (GAME RESET TO {drift[name]['actual']!r})")
            else:
                parts.append(f"{short}={value!r}")
        ws["A3"] = "Settings: " + " · ".join(parts)
        ws["A3"].font = NOTE_FONT

    runs = result.get("runs", [])
    present = [m for m in HEADLINE_METRICS if any(r["stats"].get(m) is not None for r in runs)]
    telemetry = [m for m in TELEMETRY_METRICS if any(r["stats"].get(m) is not None for r in runs)]

    header_row = 4
    columns = ["Run", "Frames", "Duration (s)", *[_label(m) for m in present], *[_label(m) for m in telemetry]]
    _write_header_row(ws, header_row, columns)

    row = header_row + 1
    for run in runs:
        stats = run["stats"]
        ws.cell(row=row, column=1, value=run["index"])
        ws.cell(row=row, column=2, value=stats.get("frames"))
        ws.cell(row=row, column=3, value=stats.get("duration_s"))
        for i, metric in enumerate(present):
            ws.cell(row=row, column=4 + i, value=stats.get(metric))
        for i, metric in enumerate(telemetry):
            ws.cell(row=row, column=4 + len(present) + i, value=stats.get(metric))
        row += 1

    last_data_row = row - 1

    # Median row, formula-driven so it follows any edits to the runs above.
    if runs:
        ws.cell(row=row, column=1, value="Median").font = Font(name=FONT, bold=True, size=10)
        for col in range(2, len(columns) + 1):
            letter = get_column_letter(col)
            ws.cell(
                row=row,
                column=col,
                value=f"=IFERROR(MEDIAN({letter}{header_row + 1}:{letter}{last_data_row}),\"\")",
            ).font = Font(name=FONT, bold=True, size=10)
        median_row = row
        row += 1
    else:
        median_row = None

    if runs:
        _style_body(ws, header_row + 1, median_row or last_data_row, 1, len(columns))
        ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

    # Bar chart: the three headline FPS figures, per run.
    charted = [m for m in CHARTED_METRICS if m in present]
    if runs and charted:
        first = 4 + present.index(charted[0])
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = f"{result['name']} — FPS per run"
        chart.y_axis.title = "FPS"
        chart.x_axis.title = "Run"
        chart.height = 8
        chart.width = 16
        data = Reference(
            ws, min_col=first, max_col=first + len(charted) - 1,
            min_row=header_row, max_row=last_data_row,
        )
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_data_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"A{row + 2}")

    # Timeline block, placed to the right so it never collides with the table.
    timeline_col = len(columns) + 2
    _write_timeline(ws, result, timeline_col, header_row, row + 2)

    _autosize(ws, len(columns))


def _write_timeline(
    ws: Worksheet, result: dict[str, Any], start_col: int, header_row: int, chart_anchor_row: int
) -> None:
    """Write per-second FPS for each run side by side, then chart it."""
    runs = [r for r in result.get("runs", []) if r.get("timeline")]
    if not runs:
        return

    label = ws.cell(row=header_row - 1, column=start_col, value="FPS over time (1-second averages)")
    label.font = SECTION_FONT

    headers = ["Second", *[f"Run {r['index']}" for r in runs]]
    _write_header_row(ws, header_row, headers, start_col=start_col)

    # Runs may differ in length; index by second so columns stay aligned.
    by_second: dict[int, dict[int, float]] = {}
    for run in runs:
        for second, fps in run["timeline"]:
            by_second.setdefault(second, {})[run["index"]] = fps

    row = header_row + 1
    for second in sorted(by_second):
        ws.cell(row=row, column=start_col, value=second)
        for i, run in enumerate(runs):
            value = by_second[second].get(run["index"])
            if value is not None:
                ws.cell(row=row, column=start_col + 1 + i, value=value)
        row += 1
    last_row = row - 1

    _style_body(ws, header_row + 1, last_row, start_col, start_col + len(runs))

    chart = LineChart()
    chart.style = 12
    chart.title = f"{result['name']} — frame rate over time"
    chart.y_axis.title = "FPS"
    chart.x_axis.title = "Seconds into capture"
    chart.height = 9
    chart.width = 24
    data = Reference(
        ws, min_col=start_col + 1, max_col=start_col + len(runs),
        min_row=header_row, max_row=last_row,
    )
    cats = Reference(ws, min_col=start_col, min_row=header_row + 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    for series in chart.series:
        series.smooth = False
        series.graphicalProperties.line.width = 12000
    ws.add_chart(chart, f"{get_column_letter(start_col)}{chart_anchor_row}")


def _build_synthetics(ws: Worksheet, session: dict[str, Any]) -> None:
    results = [r for r in session["results"] if r.get("kind") == "synthetic"]
    if not results:
        return

    ws["A1"] = "Synthetic Benchmarks"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Command-line CPU and GPU benchmarks. Median across runs."
    ws["A2"].font = NOTE_FONT

    metric_names: list[str] = []
    for result in results:
        for run in result.get("runs", []):
            for key in run.get("metrics", {}):
                if key not in metric_names:
                    metric_names.append(key)

    header_row = 4
    columns = ["Benchmark", "Runs", *metric_names, "Avg Duration (s)", "Notes"]
    _write_header_row(ws, header_row, columns)

    row = header_row + 1
    for result in results:
        runs = result.get("runs", [])
        ws.cell(row=row, column=1, value=result["name"])
        ws.cell(row=row, column=2, value=len(runs))
        for i, metric in enumerate(metric_names):
            values = sorted(
                r["metrics"][metric] for r in runs if metric in r.get("metrics", {})
            )
            if values:
                mid = len(values) // 2
                median = (
                    values[mid]
                    if len(values) % 2
                    else (values[mid - 1] + values[mid]) / 2
                )
                ws.cell(row=row, column=3 + i, value=round(median, 2))
        durations = [r.get("duration_s", 0) for r in runs]
        if durations:
            ws.cell(
                row=row, column=3 + len(metric_names),
                value=round(sum(durations) / len(durations), 1),
            )
        note = result.get("notes", "")
        failures = [r for r in runs if r.get("returncode", 0) != 0]
        if failures:
            note = (note + " " if note else "") + f"({len(failures)} run(s) exited non-zero)"
        ws.cell(row=row, column=4 + len(metric_names), value=note)
        row += 1

    last_row = row - 1
    _style_body(ws, header_row + 1, last_row, 1, len(columns))

    if metric_names and last_row > header_row:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Synthetic benchmark scores"
        chart.y_axis.title = "Score"
        chart.height = 9
        chart.width = 20
        data = Reference(
            ws, min_col=3, max_col=2 + len(metric_names),
            min_row=header_row, max_row=last_row,
        )
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"A{last_row + 3}")

    _autosize(ws, len(columns))


# --------------------------------------------------------------------------
# Entry points
# --------------------------------------------------------------------------


def build_workbook(session: dict[str, Any], out_path: str | Path) -> Path:
    """Render a session dict to an .xlsx file."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    used_names: set[str] = set()

    summary = wb.active
    summary.title = _sanitize_sheet_name("Summary", used_names)
    _build_summary(summary, session)

    env_ws = wb.create_sheet(_sanitize_sheet_name("Environment", used_names))
    _build_environment(env_ws, session)

    if any(r.get("kind") == "synthetic" for r in session["results"]):
        syn_ws = wb.create_sheet(_sanitize_sheet_name("Synthetics", used_names))
        _build_synthetics(syn_ws, session)

    for result in session["results"]:
        if result.get("kind") != "game" or not result.get("runs"):
            continue
        ws = wb.create_sheet(_sanitize_sheet_name(result["name"], used_names))
        _build_game_sheet(ws, result)

    wb.save(out_path)
    return out_path


def build_from_session_dir(session_dir: str | Path, out_path: str | Path | None = None) -> Path:
    """Rebuild the workbook from a session directory on disk.

    Reporting is deliberately separate from running: raw logs are written as the
    session goes, so a crashed or interrupted session can still be turned into a
    workbook afterwards.
    """
    session_dir = Path(session_dir)
    manifest = session_dir / "session.json"
    if not manifest.is_file():
        raise FileNotFoundError(f"No session.json in {session_dir}")
    session = json.loads(manifest.read_text(encoding="utf-8"))
    if out_path is None:
        out_path = session_dir / f"benchmarks_{session.get('session_id', 'results')}.xlsx"
    return build_workbook(session, out_path)
